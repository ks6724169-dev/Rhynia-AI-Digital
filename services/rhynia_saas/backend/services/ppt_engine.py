"""
Rhynia Intelligence SaaS — Professional AI Presentation & PPT Generator Engine
Leverages python-pptx to produce 100% native, editable Microsoft PowerPoint (.pptx) presentations
with native Excel charts, styled data tables, vector process flowcharts, 35+ executive themes,
running headers/footers, and Wikipedia/HD visual image embedding.
"""

import hashlib
import os
import uuid
import logging
from pathlib import Path
from typing import Dict, List, Any, Optional
import re
import httpx

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION

logger = logging.getLogger("rhynia.ppt_engine")

# 16:9 Widescreen dimensions
SLIDE_WIDTH_INCHES = 13.333
SLIDE_HEIGHT_INCHES = 7.5

# =========================================================
# 35+ EXECUTIVE THEMES CATALOGUE (CATEGORIZED BY DOMAIN)
# =========================================================
THEMES: Dict[str, Dict[str, Any]] = {
    # 1. Tech & AI
    "cyber_dark": {
        "label": "Cyber Dark",
        "category": "Tech & AI",
        "bg": RGBColor(11, 14, 20),           # Deep Obsidian #0B0E14
        "card_bg": RGBColor(19, 27, 38),      # Slate Obsidian #131B26
        "card_border": RGBColor(30, 58, 80),
        "text_primary": RGBColor(248, 250, 252),
        "text_secondary": RGBColor(148, 163, 184),
        "accent": RGBColor(0, 164, 239),      # Neon Cyan #00A4EF
        "accent2": RGBColor(52, 211, 153),
        "accent3": RGBColor(244, 63, 94),
        "table_header": RGBColor(14, 116, 144),
        "table_row_alt": RGBColor(15, 23, 42),
    },
    "quantum_violet": {
        "label": "Quantum Violet",
        "category": "Tech & AI",
        "bg": RGBColor(15, 10, 28),           # Deep Violet #0F0A1C
        "card_bg": RGBColor(27, 18, 51),
        "card_border": RGBColor(60, 40, 110),
        "text_primary": RGBColor(250, 245, 255),
        "text_secondary": RGBColor(192, 175, 224),
        "accent": RGBColor(168, 85, 247),     # Laser Violet #A855F7
        "accent2": RGBColor(236, 72, 153),
        "accent3": RGBColor(56, 189, 248),
        "table_header": RGBColor(88, 28, 135),
        "table_row_alt": RGBColor(24, 15, 45),
    },
    "circuit_green": {
        "label": "Circuit Green",
        "category": "Tech & AI",
        "bg": RGBColor(8, 20, 15),
        "card_bg": RGBColor(14, 36, 26),
        "card_border": RGBColor(24, 64, 46),
        "text_primary": RGBColor(240, 253, 244),
        "text_secondary": RGBColor(134, 239, 172),
        "accent": RGBColor(34, 197, 94),      # Terminal Green #22C55E
        "accent2": RGBColor(56, 189, 248),
        "accent3": RGBColor(250, 204, 21),
        "table_header": RGBColor(20, 83, 45),
        "table_row_alt": RGBColor(12, 30, 22),
    },
    "cloud_slate": {
        "label": "Cloud Slate",
        "category": "Tech & AI",
        "bg": RGBColor(15, 23, 42),          # Dark Slate #0F172A
        "card_bg": RGBColor(30, 41, 59),
        "card_border": RGBColor(51, 65, 85),
        "text_primary": RGBColor(248, 250, 252),
        "text_secondary": RGBColor(148, 163, 184),
        "accent": RGBColor(56, 189, 248),    # Sky Blue #38BDF8
        "accent2": RGBColor(99, 102, 241),
        "accent3": RGBColor(16, 185, 129),
        "table_header": RGBColor(30, 58, 138),
        "table_row_alt": RGBColor(23, 37, 84),
    },
    "devops_charcoal": {
        "label": "DevOps Charcoal",
        "category": "Tech & AI",
        "bg": RGBColor(24, 24, 27),
        "card_bg": RGBColor(39, 39, 42),
        "card_border": RGBColor(63, 63, 70),
        "text_primary": RGBColor(250, 250, 250),
        "text_secondary": RGBColor(161, 161, 170),
        "accent": RGBColor(249, 115, 22),    # Docker/GitLab Orange #F97316
        "accent2": RGBColor(14, 165, 233),
        "accent3": RGBColor(34, 197, 94),
        "table_header": RGBColor(154, 52, 18),
        "table_row_alt": RGBColor(30, 30, 34),
    },
    "ai_silicon": {
        "label": "AI Silicon",
        "category": "Tech & AI",
        "bg": RGBColor(18, 20, 24),
        "card_bg": RGBColor(28, 32, 38),
        "card_border": RGBColor(48, 54, 64),
        "text_primary": RGBColor(243, 244, 246),
        "text_secondary": RGBColor(156, 163, 175),
        "accent": RGBColor(118, 185, 0),     # Nvidia Emerald #76B900
        "accent2": RGBColor(56, 189, 248),
        "accent3": RGBColor(251, 146, 60),
        "table_header": RGBColor(40, 80, 20),
        "table_row_alt": RGBColor(22, 26, 30),
    },

    # 2. Business & Corporate
    "wall_street_navy": {
        "label": "Wall Street Navy",
        "category": "Business & Corporate",
        "bg": RGBColor(10, 22, 40),          # Royal Navy #0A1628
        "card_bg": RGBColor(16, 36, 64),
        "card_border": RGBColor(30, 60, 100),
        "text_primary": RGBColor(248, 250, 252),
        "text_secondary": RGBColor(160, 185, 215),
        "accent": RGBColor(234, 179, 8),     # Investment Gold #EAB308
        "accent2": RGBColor(56, 189, 248),
        "accent3": RGBColor(34, 197, 94),
        "table_header": RGBColor(22, 48, 86),
        "table_row_alt": RGBColor(14, 30, 54),
    },
    "corporate_azure": {
        "label": "Corporate Azure",
        "category": "Business & Corporate",
        "bg": RGBColor(12, 26, 46),
        "card_bg": RGBColor(17, 35, 61),
        "card_border": RGBColor(32, 58, 96),
        "text_primary": RGBColor(255, 255, 255),
        "text_secondary": RGBColor(165, 180, 205),
        "accent": RGBColor(0, 120, 212),     # Fluent Azure #0078D4
        "accent2": RGBColor(76, 194, 255),
        "accent3": RGBColor(16, 185, 129),
        "table_header": RGBColor(0, 120, 212),
        "table_row_alt": RGBColor(15, 30, 50),
    },
    "enterprise_gray": {
        "label": "Enterprise Gray",
        "category": "Business & Corporate",
        "bg": RGBColor(248, 250, 252),       # Clean Crisp Light
        "card_bg": RGBColor(255, 255, 255),
        "card_border": RGBColor(226, 232, 240),
        "text_primary": RGBColor(15, 23, 42),
        "text_secondary": RGBColor(100, 116, 139),
        "accent": RGBColor(37, 99, 235),     # Corporate Blue #2563EB
        "accent2": RGBColor(16, 185, 129),
        "accent3": RGBColor(249, 115, 22),
        "table_header": RGBColor(37, 99, 235),
        "table_row_alt": RGBColor(241, 245, 249),
    },
    "venture_capital": {
        "label": "Venture Capital",
        "category": "Business & Corporate",
        "bg": RGBColor(15, 28, 28),
        "card_bg": RGBColor(24, 46, 46),
        "card_border": RGBColor(40, 75, 75),
        "text_primary": RGBColor(240, 253, 250),
        "text_secondary": RGBColor(153, 246, 228),
        "accent": RGBColor(20, 184, 166),    # Deep Teal #14B8A6
        "accent2": RGBColor(245, 158, 11),
        "accent3": RGBColor(56, 189, 248),
        "table_header": RGBColor(17, 94, 89),
        "table_row_alt": RGBColor(20, 38, 38),
    },
    "strategic_monochrome": {
        "label": "Strategic Monochrome",
        "category": "Business & Corporate",
        "bg": RGBColor(18, 18, 18),
        "card_bg": RGBColor(30, 30, 30),
        "card_border": RGBColor(50, 50, 50),
        "text_primary": RGBColor(255, 255, 255),
        "text_secondary": RGBColor(180, 180, 180),
        "accent": RGBColor(220, 220, 220),   # Crisp White/Silver
        "accent2": RGBColor(140, 140, 140),
        "accent3": RGBColor(255, 255, 255),
        "table_header": RGBColor(45, 45, 45),
        "table_row_alt": RGBColor(25, 25, 25),
    },
    "consulting_cobalt": {
        "label": "Consulting Cobalt",
        "category": "Business & Corporate",
        "bg": RGBColor(10, 25, 55),
        "card_bg": RGBColor(18, 42, 85),
        "card_border": RGBColor(30, 65, 130),
        "text_primary": RGBColor(245, 250, 255),
        "text_secondary": RGBColor(170, 195, 235),
        "accent": RGBColor(30, 144, 255),    # McKinsey Cobalt #1E90FF
        "accent2": RGBColor(0, 206, 209),
        "accent3": RGBColor(255, 165, 0),
        "table_header": RGBColor(15, 50, 110),
        "table_row_alt": RGBColor(14, 32, 68),
    },
    "executive_platinum": {
        "label": "Executive Platinum",
        "category": "Business & Corporate",
        "bg": RGBColor(245, 247, 250),
        "card_bg": RGBColor(255, 255, 255),
        "card_border": RGBColor(215, 222, 232),
        "text_primary": RGBColor(20, 28, 40),
        "text_secondary": RGBColor(85, 98, 118),
        "accent": RGBColor(45, 85, 155),
        "accent2": RGBColor(30, 150, 110),
        "accent3": RGBColor(210, 110, 40),
        "table_header": RGBColor(45, 85, 155),
        "table_row_alt": RGBColor(238, 242, 248),
    },

    # 3. Startups & Pitch Decks
    "yc_orange": {
        "label": "Y-Combinator Orange",
        "category": "Startups & Pitch Decks",
        "bg": RGBColor(15, 15, 17),
        "card_bg": RGBColor(26, 26, 30),
        "card_border": RGBColor(50, 48, 56),
        "text_primary": RGBColor(255, 255, 255),
        "text_secondary": RGBColor(175, 175, 185),
        "accent": RGBColor(255, 102, 0),     # YC Orange #FF6600
        "accent2": RGBColor(251, 146, 60),
        "accent3": RGBColor(34, 197, 94),
        "table_header": RGBColor(180, 70, 0),
        "table_row_alt": RGBColor(22, 22, 25),
    },
    "unicorn_purple": {
        "label": "Unicorn Purple",
        "category": "Startups & Pitch Decks",
        "bg": RGBColor(16, 12, 32),
        "card_bg": RGBColor(28, 22, 54),
        "card_border": RGBColor(55, 44, 98),
        "text_primary": RGBColor(255, 255, 255),
        "text_secondary": RGBColor(195, 185, 225),
        "accent": RGBColor(139, 92, 246),    # Silicon Purple #8B5CF6
        "accent2": RGBColor(244, 63, 94),
        "accent3": RGBColor(56, 189, 248),
        "table_header": RGBColor(91, 33, 182),
        "table_row_alt": RGBColor(22, 17, 42),
    },
    "saas_indigo": {
        "label": "SaaS Indigo",
        "category": "Startups & Pitch Decks",
        "bg": RGBColor(12, 15, 30),
        "card_bg": RGBColor(21, 26, 52),
        "card_border": RGBColor(40, 50, 95),
        "text_primary": RGBColor(245, 247, 255),
        "text_secondary": RGBColor(160, 175, 215),
        "accent": RGBColor(99, 102, 241),    # Stripe Indigo #6366F1
        "accent2": RGBColor(168, 85, 247),
        "accent3": RGBColor(16, 185, 129),
        "table_header": RGBColor(55, 48, 163),
        "table_row_alt": RGBColor(17, 21, 42),
    },
    "pitch_dark": {
        "label": "Pitch Night Dark",
        "category": "Startups & Pitch Decks",
        "bg": RGBColor(8, 8, 10),
        "card_bg": RGBColor(18, 18, 22),
        "card_border": RGBColor(38, 38, 45),
        "text_primary": RGBColor(255, 255, 255),
        "text_secondary": RGBColor(160, 160, 170),
        "accent": RGBColor(244, 63, 94),     # Vibrant Rose #F43F5E
        "accent2": RGBColor(56, 189, 248),
        "accent3": RGBColor(250, 204, 21),
        "table_header": RGBColor(159, 18, 57),
        "table_row_alt": RGBColor(14, 14, 17),
    },
    "bootstrapped_green": {
        "label": "Bootstrapped Green",
        "category": "Startups & Pitch Decks",
        "bg": RGBColor(10, 24, 18),
        "card_bg": RGBColor(18, 42, 32),
        "card_border": RGBColor(32, 72, 55),
        "text_primary": RGBColor(240, 253, 244),
        "text_secondary": RGBColor(155, 225, 185),
        "accent": RGBColor(16, 185, 129),    # Profit Green #10B981
        "accent2": RGBColor(245, 158, 11),
        "accent3": RGBColor(56, 189, 248),
        "table_header": RGBColor(6, 95, 70),
        "table_row_alt": RGBColor(14, 32, 25),
    },
    "founder_minimal": {
        "label": "Founder Minimal",
        "category": "Startups & Pitch Decks",
        "bg": RGBColor(255, 255, 255),
        "card_bg": RGBColor(245, 245, 247),
        "card_border": RGBColor(220, 220, 225),
        "text_primary": RGBColor(10, 10, 15),
        "text_secondary": RGBColor(90, 90, 100),
        "accent": RGBColor(0, 0, 0),         # Pure Black Bold
        "accent2": RGBColor(0, 120, 212),
        "accent3": RGBColor(220, 38, 38),
        "table_header": RGBColor(20, 20, 25),
        "table_row_alt": RGBColor(240, 240, 244),
    },

    # 4. Education, Science & Research
    "academic_slate": {
        "label": "Academic Slate",
        "category": "Education & Science",
        "bg": RGBColor(18, 24, 38),
        "card_bg": RGBColor(28, 38, 58),
        "card_border": RGBColor(45, 60, 90),
        "text_primary": RGBColor(245, 248, 255),
        "text_secondary": RGBColor(165, 180, 210),
        "accent": RGBColor(79, 145, 255),
        "accent2": RGBColor(46, 204, 113),
        "accent3": RGBColor(241, 196, 15),
        "table_header": RGBColor(25, 45, 75),
        "table_row_alt": RGBColor(22, 30, 48),
    },
    "university_burgundy": {
        "label": "University Burgundy",
        "category": "Education & Science",
        "bg": RGBColor(28, 12, 18),
        "card_bg": RGBColor(46, 20, 30),
        "card_border": RGBColor(80, 35, 52),
        "text_primary": RGBColor(255, 245, 248),
        "text_secondary": RGBColor(215, 175, 190),
        "accent": RGBColor(225, 29, 72),     # Harvard/Oxford Crimson
        "accent2": RGBColor(234, 179, 8),
        "accent3": RGBColor(56, 189, 248),
        "table_header": RGBColor(136, 19, 55),
        "table_row_alt": RGBColor(36, 15, 24),
    },
    "research_teal": {
        "label": "Research Teal",
        "category": "Education & Science",
        "bg": RGBColor(10, 28, 30),
        "card_bg": RGBColor(18, 45, 48),
        "card_border": RGBColor(30, 72, 76),
        "text_primary": RGBColor(240, 253, 250),
        "text_secondary": RGBColor(150, 215, 210),
        "accent": RGBColor(45, 212, 191),    # Scientific Aqua #2DD4BF
        "accent2": RGBColor(96, 165, 250),
        "accent3": RGBColor(244, 114, 182),
        "table_header": RGBColor(19, 78, 74),
        "table_row_alt": RGBColor(14, 36, 38),
    },
    "chalkboard_dark": {
        "label": "Chalkboard Dark",
        "category": "Education & Science",
        "bg": RGBColor(18, 32, 26),
        "card_bg": RGBColor(28, 48, 40),
        "card_border": RGBColor(45, 75, 62),
        "text_primary": RGBColor(255, 255, 255),
        "text_secondary": RGBColor(180, 210, 195),
        "accent": RGBColor(253, 224, 71),    # Chalk Yellow #FDE047
        "accent2": RGBColor(110, 231, 183),
        "accent3": RGBColor(147, 197, 253),
        "table_header": RGBColor(20, 55, 42),
        "table_row_alt": RGBColor(22, 38, 32),
    },
    "scholar_paper": {
        "label": "Scholar Paper",
        "category": "Education & Science",
        "bg": RGBColor(250, 248, 245),       # Soft Paper Canvas
        "card_bg": RGBColor(255, 255, 255),
        "card_border": RGBColor(225, 220, 212),
        "text_primary": RGBColor(30, 25, 20),
        "text_secondary": RGBColor(95, 85, 75),
        "accent": RGBColor(160, 60, 40),     # Terracotta #A03C28
        "accent2": RGBColor(40, 110, 130),
        "accent3": RGBColor(180, 140, 40),
        "table_header": RGBColor(140, 50, 30),
        "table_row_alt": RGBColor(242, 238, 232),
    },
    "stem_cyan": {
        "label": "STEM Cyan",
        "category": "Education & Science",
        "bg": RGBColor(12, 22, 34),
        "card_bg": RGBColor(20, 36, 56),
        "card_border": RGBColor(35, 60, 92),
        "text_primary": RGBColor(245, 250, 255),
        "text_secondary": RGBColor(160, 190, 220),
        "accent": RGBColor(6, 182, 212),     # Electric Cyan #06B6D4
        "accent2": RGBColor(59, 130, 246),
        "accent3": RGBColor(16, 185, 129),
        "table_header": RGBColor(14, 116, 144),
        "table_row_alt": RGBColor(16, 28, 44),
    },

    # 5. Creative, Marketing & Media
    "sunset_coral": {
        "label": "Sunset Coral",
        "category": "Creative & Media",
        "bg": RGBColor(28, 14, 18),
        "card_bg": RGBColor(48, 24, 30),
        "card_border": RGBColor(82, 42, 52),
        "text_primary": RGBColor(255, 245, 245),
        "text_secondary": RGBColor(225, 180, 185),
        "accent": RGBColor(251, 113, 133),   # Warm Coral #FB7185
        "accent2": RGBColor(251, 146, 60),
        "accent3": RGBColor(250, 204, 21),
        "table_header": RGBColor(190, 24, 93),
        "table_row_alt": RGBColor(38, 18, 24),
    },
    "neon_pulse": {
        "label": "Neon Pulse",
        "category": "Creative & Media",
        "bg": RGBColor(12, 10, 22),
        "card_bg": RGBColor(24, 18, 42),
        "card_border": RGBColor(52, 38, 90),
        "text_primary": RGBColor(255, 255, 255),
        "text_secondary": RGBColor(200, 185, 230),
        "accent": RGBColor(244, 63, 94),     # Hot Pink / Neon
        "accent2": RGBColor(56, 189, 248),
        "accent3": RGBColor(168, 85, 247),
        "table_header": RGBColor(159, 18, 57),
        "table_row_alt": RGBColor(18, 14, 32),
    },
    "minimal_peach": {
        "label": "Minimal Peach",
        "category": "Creative & Media",
        "bg": RGBColor(253, 248, 245),
        "card_bg": RGBColor(255, 255, 255),
        "card_border": RGBColor(240, 224, 216),
        "text_primary": RGBColor(35, 25, 25),
        "text_secondary": RGBColor(115, 95, 95),
        "accent": RGBColor(249, 115, 22),    # Warm Peach #F97316
        "accent2": RGBColor(236, 72, 153),
        "accent3": RGBColor(14, 165, 233),
        "table_header": RGBColor(194, 65, 12),
        "table_row_alt": RGBColor(248, 240, 235),
    },
    "bold_crimson": {
        "label": "Bold Crimson",
        "category": "Creative & Media",
        "bg": RGBColor(18, 10, 12),
        "card_bg": RGBColor(34, 16, 20),
        "card_border": RGBColor(65, 30, 38),
        "text_primary": RGBColor(255, 255, 255),
        "text_secondary": RGBColor(220, 175, 180),
        "accent": RGBColor(225, 29, 72),     # High Voltage Crimson #E11D48
        "accent2": RGBColor(251, 146, 60),
        "accent3": RGBColor(250, 204, 21),
        "table_header": RGBColor(159, 18, 57),
        "table_row_alt": RGBColor(26, 12, 16),
    },
    "pastel_creative": {
        "label": "Pastel Creative",
        "category": "Creative & Media",
        "bg": RGBColor(246, 248, 252),
        "card_bg": RGBColor(255, 255, 255),
        "card_border": RGBColor(225, 232, 245),
        "text_primary": RGBColor(25, 30, 45),
        "text_secondary": RGBColor(95, 105, 130),
        "accent": RGBColor(129, 140, 248),   # Pastel Indigo #818CF8
        "accent2": RGBColor(244, 114, 182),
        "accent3": RGBColor(52, 211, 153),
        "table_header": RGBColor(79, 70, 229),
        "table_row_alt": RGBColor(240, 244, 250),
    },
    "editorial_beige": {
        "label": "Editorial Beige",
        "category": "Creative & Media",
        "bg": RGBColor(247, 244, 238),
        "card_bg": RGBColor(255, 255, 255),
        "card_border": RGBColor(225, 220, 208),
        "text_primary": RGBColor(24, 22, 18),
        "text_secondary": RGBColor(90, 85, 75),
        "accent": RGBColor(140, 105, 65),    # Warm Bronze #8C6941
        "accent2": RGBColor(60, 95, 80),
        "accent3": RGBColor(165, 75, 60),
        "table_header": RGBColor(110, 80, 45),
        "table_row_alt": RGBColor(240, 236, 228),
    },
    "vibrant_tropical": {
        "label": "Vibrant Tropical",
        "category": "Creative & Media",
        "bg": RGBColor(12, 20, 24),
        "card_bg": RGBColor(20, 36, 42),
        "card_border": RGBColor(35, 62, 72),
        "text_primary": RGBColor(255, 255, 255),
        "text_secondary": RGBColor(165, 210, 220),
        "accent": RGBColor(20, 184, 166),    # Tropical Teal #14B8A6
        "accent2": RGBColor(249, 115, 22),
        "accent3": RGBColor(234, 179, 8),
        "table_header": RGBColor(15, 118, 110),
        "table_row_alt": RGBColor(16, 28, 34),
    },

    # 6. Medical, Eco & Sustainability
    "clinical_blue": {
        "label": "Clinical Blue",
        "category": "Medical & Health",
        "bg": RGBColor(245, 250, 255),
        "card_bg": RGBColor(255, 255, 255),
        "card_border": RGBColor(215, 232, 245),
        "text_primary": RGBColor(15, 30, 50),
        "text_secondary": RGBColor(75, 100, 130),
        "accent": RGBColor(2, 132, 199),     # Medical Blue #0284C7
        "accent2": RGBColor(13, 148, 136),
        "accent3": RGBColor(225, 29, 72),
        "table_header": RGBColor(3, 105, 161),
        "table_row_alt": RGBColor(238, 246, 254),
    },
    "eco_forest": {
        "label": "Eco Forest",
        "category": "Sustainability",
        "bg": RGBColor(8, 22, 16),
        "card_bg": RGBColor(15, 40, 28),
        "card_border": RGBColor(28, 68, 48),
        "text_primary": RGBColor(240, 253, 244),
        "text_secondary": RGBColor(150, 220, 180),
        "accent": RGBColor(34, 197, 94),      # Forest Green #22C55E
        "accent2": RGBColor(234, 179, 8),
        "accent3": RGBColor(56, 189, 248),
        "table_header": RGBColor(20, 83, 45),
        "table_row_alt": RGBColor(12, 32, 22),
    },
    "pharma_mint": {
        "label": "Pharma Mint",
        "category": "Medical & Health",
        "bg": RGBColor(246, 253, 250),
        "card_bg": RGBColor(255, 255, 255),
        "card_border": RGBColor(215, 240, 230),
        "text_primary": RGBColor(18, 38, 30),
        "text_secondary": RGBColor(75, 115, 95),
        "accent": RGBColor(13, 148, 136),    # Mint Teal #0D9488
        "accent2": RGBColor(59, 130, 246),
        "accent3": RGBColor(244, 63, 94),
        "table_header": RGBColor(15, 118, 110),
        "table_row_alt": RGBColor(236, 249, 244),
    },
    "renewable_lime": {
        "label": "Renewable Lime",
        "category": "Sustainability",
        "bg": RGBColor(14, 24, 12),
        "card_bg": RGBColor(24, 44, 22),
        "card_border": RGBColor(45, 78, 40),
        "text_primary": RGBColor(247, 254, 231),
        "text_secondary": RGBColor(185, 225, 145),
        "accent": RGBColor(132, 204, 22),    # Solar Lime #84CC16
        "accent2": RGBColor(234, 179, 8),
        "accent3": RGBColor(14, 165, 233),
        "table_header": RGBColor(54, 83, 20),
        "table_row_alt": RGBColor(18, 34, 16),
    },
    "biotech_aqua": {
        "label": "BioTech Aqua",
        "category": "Medical & Health",
        "bg": RGBColor(10, 25, 35),
        "card_bg": RGBColor(18, 42, 58),
        "card_border": RGBColor(32, 68, 92),
        "text_primary": RGBColor(240, 250, 255),
        "text_secondary": RGBColor(155, 205, 225),
        "accent": RGBColor(6, 182, 212),     # Bio Aqua #06B6D4
        "accent2": RGBColor(16, 185, 129),
        "accent3": RGBColor(244, 63, 94),
        "table_header": RGBColor(14, 116, 144),
        "table_row_alt": RGBColor(14, 34, 46),
    },

    # Aliases for backward compatibility
    "executive_dark": {
        "label": "Executive Dark",
        "category": "Tech & AI",
        "bg": RGBColor(11, 14, 20),
        "card_bg": RGBColor(19, 27, 38),
        "card_border": RGBColor(30, 58, 80),
        "text_primary": RGBColor(248, 250, 252),
        "text_secondary": RGBColor(148, 163, 184),
        "accent": RGBColor(0, 164, 239),
        "accent2": RGBColor(52, 211, 153),
        "accent3": RGBColor(244, 63, 94),
        "table_header": RGBColor(14, 116, 144),
        "table_row_alt": RGBColor(15, 23, 42),
    },
    "emerald_minimal": {
        "label": "Emerald Minimal",
        "category": "Sustainability",
        "bg": RGBColor(8, 22, 16),
        "card_bg": RGBColor(15, 40, 28),
        "card_border": RGBColor(28, 68, 48),
        "text_primary": RGBColor(240, 253, 244),
        "text_secondary": RGBColor(150, 220, 180),
        "accent": RGBColor(34, 197, 94),
        "accent2": RGBColor(234, 179, 8),
        "accent3": RGBColor(56, 189, 248),
        "table_header": RGBColor(20, 83, 45),
        "table_row_alt": RGBColor(12, 32, 22),
    },
    "clean_light": {
        "label": "Clean Light",
        "category": "Business & Corporate",
        "bg": RGBColor(248, 250, 252),
        "card_bg": RGBColor(255, 255, 255),
        "card_border": RGBColor(226, 232, 240),
        "text_primary": RGBColor(15, 23, 42),
        "text_secondary": RGBColor(100, 116, 139),
        "accent": RGBColor(37, 99, 235),
        "accent2": RGBColor(16, 185, 129),
        "accent3": RGBColor(249, 115, 22),
        "table_header": RGBColor(37, 99, 235),
        "table_row_alt": RGBColor(241, 245, 249),
    }
}


class PPTEngine:
    """
    World-Class PowerPoint Presentation Engine producing 16:9 widescreen presentations
    with 35+ themes, headers/footers, real Excel charts, tables, diagrams, and HD images.
    """

    def __init__(self, output_dir: Optional[Path] = None):
        self.output_dir = output_dir or Path("database/generated_ppt")
        self.output_dir.mkdir(parents=True, exist_ok=True)
        self.image_cache_dir = Path("database/image_cache")
        self.image_cache_dir.mkdir(parents=True, exist_ok=True)

    def _apply_slide_background(self, slide, theme: Dict[str, Any]):
        """Sets slide solid background color."""
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = theme["bg"]

    def _add_header(self, slide, theme: Dict[str, Any], title: str, category_tag: str = "EXECUTIVE STRATEGY"):
        """Standardized slide header with category pill, title, and accent rule."""
        clean_cat = (category_tag or "STRATEGIC OVERVIEW").upper()
        # Category Pill
        tag_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.5), Inches(0.32))
        tf_tag = tag_box.text_frame
        tf_tag.word_wrap = True
        tf_tag.margin_left = tf_tag.margin_top = tf_tag.margin_right = tf_tag.margin_bottom = 0
        p_tag = tf_tag.paragraphs[0]
        p_tag.text = f"❖  {clean_cat}"
        p_tag.font.bold = True
        p_tag.font.size = Pt(11)
        p_tag.font.color.rgb = theme["accent"]

        # Main Title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.72), Inches(11.5), Inches(0.65))
        tf_title = title_box.text_frame
        tf_title.word_wrap = True
        tf_title.margin_left = tf_title.margin_top = tf_title.margin_right = tf_title.margin_bottom = 0
        p_title = tf_title.paragraphs[0]
        p_title.text = title or "Executive Overview"
        p_title.font.bold = True
        p_title.font.size = Pt(24)
        p_title.font.color.rgb = theme["text_primary"]

        # Header accent separator line
        accent_line = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(0.8), Inches(1.42), Inches(11.733), Inches(0.02)
        )
        accent_line.fill.solid()
        accent_line.fill.fore_color.rgb = theme.get("card_border", theme["accent"])
        accent_line.line.fill.background()

    def _add_footer(self, slide, theme: Dict[str, Any], current_slide: int, total_slides: int, presenter: str = "Rhynia AI"):
        """Standardized professional running footer with presenter name, confidentiality, and slide number."""
        footer_y = Inches(6.95)
        footer_w = Inches(11.733)

        # Top border rule for footer
        f_line = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(0.8), footer_y, footer_w, Inches(0.015)
        )
        f_line.fill.solid()
        f_line.fill.fore_color.rgb = theme.get("card_border", theme["accent"])
        f_line.line.fill.background()

        # Left: Presenter & Brand
        box_left = slide.shapes.add_textbox(Inches(0.8), footer_y + Inches(0.06), Inches(4.5), Inches(0.3))
        tf_l = box_left.text_frame
        tf_l.word_wrap = False
        tf_l.margin_left = tf_l.margin_top = tf_l.margin_right = tf_l.margin_bottom = 0
        p_l = tf_l.paragraphs[0]
        p_l.text = f"Presented by: {presenter} | Rhynia AI"
        p_l.font.size = Pt(9.5)
        p_l.font.color.rgb = theme["text_secondary"]

        # Center: Confidentiality / Date
        box_mid = slide.shapes.add_textbox(Inches(5.5), footer_y + Inches(0.06), Inches(3.5), Inches(0.3))
        tf_m = box_mid.text_frame
        tf_m.word_wrap = False
        tf_m.margin_left = tf_m.margin_top = tf_m.margin_right = tf_m.margin_bottom = 0
        p_m = tf_m.paragraphs[0]
        p_m.text = "Confidential • Executive Edition • 2026"
        p_m.alignment = PP_ALIGN.CENTER
        p_m.font.size = Pt(9.5)
        p_m.font.color.rgb = theme["text_secondary"]

        # Right: Slide Number
        box_right = slide.shapes.add_textbox(Inches(10.5), footer_y + Inches(0.06), Inches(2.033), Inches(0.3))
        tf_r = box_right.text_frame
        tf_r.word_wrap = False
        tf_r.margin_left = tf_r.margin_top = tf_r.margin_right = tf_r.margin_bottom = 0
        p_r = tf_r.paragraphs[0]
        p_r.text = f"Slide {current_slide:02d} / {total_slides:02d}"
        p_r.alignment = PP_ALIGN.RIGHT
        p_r.font.bold = True
        p_r.font.size = Pt(9.5)
        p_r.font.color.rgb = theme["accent"]

    def _download_and_cache_image(self, url: str) -> Optional[Path]:
        """Downloads an image from URL and caches locally for native PowerPoint picture embedding."""
        if not url or not isinstance(url, str) or not url.startswith("http"):
            return None
        try:
            # Strip tracking query params
            clean_url = url.split("?")[0].strip()
            # PowerPoint requires JPG/PNG/BMP (skip raw SVG)
            if clean_url.lower().endswith(".svg"):
                return None

            url_hash = hashlib.md5(clean_url.encode("utf-8")).hexdigest()
            ext = ".png" if clean_url.lower().endswith(".png") else ".jpg"
            target_path = self.image_cache_dir / f"{url_hash}{ext}"
            if target_path.exists() and target_path.stat().st_size > 1500:
                return target_path

            headers = {
                "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36 RhyniaPPT/1.0",
                "Referer": "https://commons.wikimedia.org/"
            }
            with httpx.Client(timeout=10.0, follow_redirects=True) as client:
                res = client.get(clean_url, headers=headers)
                if res.status_code == 200 and len(res.content) > 1500:
                    with open(target_path, "wb") as f:
                        f.write(res.content)
                    return target_path
        except Exception as e:
            logger.warning(f"Failed to cache slide image from {url}: {e}")
        return None

    # ==========================================
    # SLIDE BUILDER 1: TITLE & COVER SLIDE
    # ==========================================
    # ==========================================
    # SLIDE BUILDER 1: TITLE & COVER SLIDE (EXECUTIVE INTRO ARCHITECTURE)
    # ==========================================
    def add_title_slide(self, prs, theme: Dict[str, Any], data: Dict[str, Any]):
        blank_slide_layout = prs.slide_layouts[6]
        slide = prs.slides.add_slide(blank_slide_layout)
        self._apply_slide_background(slide, theme)

        title = data.get("title", "Executive Presentation")
        subtitle = data.get("subtitle", "Strategic Roadmap & High-Impact Analysis")
        category = (data.get("category", "EXECUTIVE STRATEGY")).upper()
        presenter = data.get("presented_by") or data.get("presenter", "Rhynia Intelligence")
        date_str = data.get("date", "2026 Edition")
        location = data.get("location", "Corporate HQ / New Delhi")
        goal = data.get("goal", "Strategic Implementation & Action")

        # Top Accent Header Stripe
        top_bar = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(0), Inches(0), Inches(SLIDE_WIDTH_INCHES), Inches(0.08)
        )
        top_bar.fill.solid()
        top_bar.fill.fore_color.rgb = theme["accent"]
        top_bar.line.fill.background()

        # Check if hero image provided
        image_url = data.get("image_url")
        image_path = self._download_and_cache_image(image_url) if image_url else None

        if image_path:
            # Title Box (Left Side)
            title_box_shape = slide.shapes.add_shape(
                MSO_SHAPE.ROUNDED_RECTANGLE,
                Inches(0.8), Inches(0.85), Inches(7.6), Inches(3.5)
            )
            title_box_shape.fill.solid()
            title_box_shape.fill.fore_color.rgb = theme["card_bg"]
            title_box_shape.line.color.rgb = theme.get("card_border", theme["accent"])
            title_box_shape.line.width = Pt(1.5)

            # Category inside Box
            cat_box = slide.shapes.add_textbox(Inches(1.0), Inches(1.05), Inches(7.2), Inches(0.4))
            tf_c = cat_box.text_frame
            p_c = tf_c.paragraphs[0]
            p_c.text = f"❖  {category}  ❖"
            p_c.font.bold = True
            p_c.font.size = Pt(11)
            p_c.font.color.rgb = theme["accent"]

            # Main Title inside Box
            title_box = slide.shapes.add_textbox(Inches(1.0), Inches(1.5), Inches(7.2), Inches(1.6))
            tf_t = title_box.text_frame
            tf_t.word_wrap = True
            p_t = tf_t.paragraphs[0]
            p_t.text = title
            p_t.font.bold = True
            p_t.font.size = Pt(32)
            p_t.font.color.rgb = theme["text_primary"]

            # Subtitle inside Box
            sub_box = slide.shapes.add_textbox(Inches(1.0), Inches(3.15), Inches(7.2), Inches(0.9))
            tf_s = sub_box.text_frame
            tf_s.word_wrap = True
            p_s = tf_s.paragraphs[0]
            p_s.text = subtitle
            p_s.font.size = Pt(13)
            p_s.font.color.rgb = theme["text_secondary"]

            # Right Hero Image Card
            try:
                img_card = slide.shapes.add_shape(
                    MSO_SHAPE.ROUNDED_RECTANGLE,
                    Inches(8.7), Inches(0.85), Inches(3.8), Inches(3.5)
                )
                img_card.fill.solid()
                img_card.fill.fore_color.rgb = theme["card_bg"]
                img_card.line.color.rgb = theme.get("card_border", theme["accent"])
                img_card.line.width = Pt(1.5)

                slide.shapes.add_picture(str(image_path), Inches(8.8), Inches(0.95), Inches(3.6), Inches(3.3))
            except Exception as e:
                logger.warning(f"Error embedding title image: {e}")
        else:
            # Full Centered Title Box Frame
            title_box_shape = slide.shapes.add_shape(
                MSO_SHAPE.ROUNDED_RECTANGLE,
                Inches(1.5), Inches(0.85), Inches(10.333), Inches(3.5)
            )
            title_box_shape.fill.solid()
            title_box_shape.fill.fore_color.rgb = theme["card_bg"]
            title_box_shape.line.color.rgb = theme.get("card_border", theme["accent"])
            title_box_shape.line.width = Pt(1.5)

            # Category inside Centered Box
            cat_box = slide.shapes.add_textbox(Inches(1.8), Inches(1.1), Inches(9.733), Inches(0.4))
            tf_c = cat_box.text_frame
            p_c = tf_c.paragraphs[0]
            p_c.alignment = PP_ALIGN.CENTER
            p_c.text = f"❖  PROJECT NATURE • {category}  ❖"
            p_c.font.bold = True
            p_c.font.size = Pt(12)
            p_c.font.color.rgb = theme["accent"]

            # Main Title inside Centered Box
            title_box = slide.shapes.add_textbox(Inches(1.8), Inches(1.6), Inches(9.733), Inches(1.6))
            tf_t = title_box.text_frame
            tf_t.word_wrap = True
            p_t = tf_t.paragraphs[0]
            p_t.alignment = PP_ALIGN.CENTER
            p_t.text = title
            p_t.font.bold = True
            p_t.font.size = Pt(38)
            p_t.font.color.rgb = theme["text_primary"]

            # Subtitle inside Centered Box
            sub_box = slide.shapes.add_textbox(Inches(1.8), Inches(3.2), Inches(9.733), Inches(0.9))
            tf_s = sub_box.text_frame
            tf_s.word_wrap = True
            p_s = tf_s.paragraphs[0]
            p_s.alignment = PP_ALIGN.CENTER
            p_s.text = subtitle
            p_s.font.size = Pt(14.5)
            p_s.font.color.rgb = theme["text_secondary"]

        # Decorative Divider Motif (Double Ring Emblem / Center Motif)
        motif_box = slide.shapes.add_textbox(Inches(1.5), Inches(4.45), Inches(10.333), Inches(0.35))
        tf_m = motif_box.text_frame
        p_m = tf_m.paragraphs[0]
        p_m.alignment = PP_ALIGN.CENTER
        p_m.text = "──────────────  ◎  ──────────────"
        p_m.font.bold = True
        p_m.font.size = Pt(11)
        p_m.font.color.rgb = theme["accent"]

        # 4-Point Structured Metadata Grid (4 Sleek Rounded Cards)
        card_w = Inches(2.7)
        card_h = Inches(1.45)
        card_gap = Inches(0.24)
        start_x = Inches(0.95)
        card_y = Inches(4.9)

        meta_items = [
            ("👤  PRESENTED BY", presenter, "Rhynia Intelligence"),
            ("📅  DATE", date_str, "Executive Release"),
            ("📍  LOCATION", location, "Corporate HQ"),
            ("🎯  GOAL", goal, "Strategic Focus"),
        ]

        for i, (label, val, sub) in enumerate(meta_items):
            cx = start_x + (i * (card_w + card_gap))
            c_shape = slide.shapes.add_shape(
                MSO_SHAPE.ROUNDED_RECTANGLE,
                cx, card_y, card_w, card_h
            )
            c_shape.fill.solid()
            c_shape.fill.fore_color.rgb = theme["card_bg"]
            c_shape.line.color.rgb = theme.get("card_border", theme["accent"])
            c_shape.line.width = Pt(1)

            t_box = slide.shapes.add_textbox(cx + Inches(0.12), card_y + Inches(0.12), card_w - Inches(0.24), card_h - Inches(0.24))
            tf = t_box.text_frame
            tf.word_wrap = True
            tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

            # Label
            p_lbl = tf.paragraphs[0]
            p_lbl.text = label
            p_lbl.font.bold = True
            p_lbl.font.size = Pt(9)
            p_lbl.font.color.rgb = theme["accent"]

            # Value
            p_val = tf.add_paragraph()
            clean_val = val.strip()
            p_val.text = clean_val
            p_val.font.bold = True
            p_val.font.size = Pt(9.5) if len(clean_val) > 32 else Pt(11)
            p_val.font.color.rgb = theme["text_primary"]

            # Sub-caption
            p_sub = tf.add_paragraph()
            p_sub.text = sub
            p_sub.font.size = Pt(8.5)
            p_sub.font.color.rgb = theme["text_secondary"]

        # Bottom Running Accent Rule & Confidentiality Tag
        footer_line = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(0.8), Inches(6.8), Inches(11.733), Inches(0.015)
        )
        footer_line.fill.solid()
        footer_line.fill.fore_color.rgb = theme.get("card_border", theme["accent"])
        footer_line.line.fill.background()

        bot_box = slide.shapes.add_textbox(Inches(0.8), Inches(6.88), Inches(11.733), Inches(0.3))
        tf_b = bot_box.text_frame
        p_b = tf_b.paragraphs[0]
        p_b.alignment = PP_ALIGN.CENTER
        p_b.text = "CONFIDENTIAL • EXECUTIVE 16:9 MASTER DECK • 2026 EDITION"
        p_b.font.size = Pt(9)
        p_b.font.color.rgb = theme["text_secondary"]

    # ==========================================
    # SLIDE BUILDER 2: AGENDA SLIDE
    # ==========================================
    def add_agenda_slide(self, prs, theme: Dict[str, Any], data: Dict[str, Any], current_slide: int = 2, total_slides: int = 8, presenter: str = "Rhynia AI"):
        blank_slide_layout = prs.slide_layouts[6]
        slide = prs.slides.add_slide(blank_slide_layout)
        self._apply_slide_background(slide, theme)
        self._add_header(slide, theme, data.get("title", "Presentation Agenda"), data.get("category", "AGENDA & OVERVIEW"))
        self._add_footer(slide, theme, current_slide, total_slides, presenter)

        items = data.get("items", [])
        if not items:
            s_title = data.get("title", "मुख्य विषय")
            items = [
                {"title": f"1. {s_title} की प्रस्तावना", "desc": f"{s_title} के मौलिक सिद्धांतों, पृष्ठभूमि और प्राथमिक उद्देश्यों का परिचय।"},
                {"title": "2. संरचना एवं कार्यप्रणाली", "desc": "इसके आंतरिक घटकों, प्रणालियों और परिचालन विधियों का संपूर्ण तकनीकी विश्लेषण।"},
                {"title": "3. प्रभाव एवं सांख्यिकी", "desc": "व्यावहारिक प्रयोग, अनुभवजन्य आंकड़े और प्रमुख प्रदर्शन संकेतकों का मूल्यांकन।"},
                {"title": "4. निष्कर्ष एवं भावी दिशा", "desc": "भविष्य की रणनीतियां, सतत नवाचार और मुख्य रणनीतिक निष्कर्ष।"}
            ]

        count = len(items)
        cols = 2 if count > 3 else 1
        rows = (count + 1) // 2 if cols == 2 else count

        card_w = Inches(5.6) if cols == 2 else Inches(11.733)
        card_h = Inches(4.8 / max(rows, 1)) - Inches(0.15)
        top_start = Inches(1.7)

        for idx, item in enumerate(items):
            if isinstance(item, str):
                item = {"title": item, "desc": ""}
            elif not isinstance(item, dict):
                item = {"title": str(item), "desc": ""}
            c_idx = idx % cols
            r_idx = idx // cols
            left = Inches(0.8) + (c_idx * Inches(6.1))
            top = top_start + (r_idx * (card_h + Inches(0.18)))

            card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, card_w, card_h)
            card.fill.solid()
            card.fill.fore_color.rgb = theme["card_bg"]
            card.line.color.rgb = theme.get("card_border", theme["accent"])
            card.line.width = Pt(1)

            # Number badge
            badge = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left + Inches(0.2), top + Inches(0.18), Inches(0.55), Inches(0.55))
            badge.fill.solid()
            badge.fill.fore_color.rgb = theme["accent"]
            badge.line.fill.background()
            badge.text_frame.text = f"{idx + 1:02d}"
            p_badge = badge.text_frame.paragraphs[0]
            p_badge.font.bold = True
            p_badge.font.size = Pt(13)
            p_badge.font.color.rgb = theme["bg"]
            p_badge.alignment = PP_ALIGN.CENTER

            # Text
            text_left = left + Inches(0.9)
            text_w = card_w - Inches(1.1)
            t_box = slide.shapes.add_textbox(text_left, top + Inches(0.12), text_w, card_h - Inches(0.2))
            tf = t_box.text_frame
            tf.word_wrap = True

            p_title = tf.paragraphs[0]
            p_title.text = item.get("title", f"Agenda Item {idx + 1}")
            p_title.font.bold = True
            p_title.font.size = Pt(14)
            p_title.font.color.rgb = theme["text_primary"]
            p_title.space_after = Pt(4)

            p_desc = tf.add_paragraph()
            p_desc.text = item.get("desc", "")
            p_desc.font.size = Pt(11)
            p_desc.font.color.rgb = theme["text_secondary"]

    # ==========================================
    # SLIDE BUILDER 3: KPI & METRIC CARDS
    # ==========================================
    def add_card_content_slide(self, prs, theme: Dict[str, Any], data: Dict[str, Any], current_slide: int = 3, total_slides: int = 8, presenter: str = "Rhynia AI"):
        blank_slide_layout = prs.slide_layouts[6]
        slide = prs.slides.add_slide(blank_slide_layout)
        self._apply_slide_background(slide, theme)
        self._add_header(slide, theme, data.get("title", "Key Strategic Metrics"), data.get("category", "EXECUTIVE INSIGHTS"))
        self._add_footer(slide, theme, current_slide, total_slides, presenter)

        cards = data.get("cards", [])
        if not cards:
            s_title = data.get("title", "मुख्य बिंदु")
            cards = [
                {"title": f"{s_title} - मुख्य अवधारणा", "desc": f"{s_title} के मौलिक सिद्धांतों और प्रमुख वैज्ञानिक व व्यावहारिक पहलुओं का विस्तृत विश्लेषण।"},
                {"title": f"{s_title} - महत्व व प्रभाव", "desc": f"वर्तमान परिप्रेक्ष्य में इसके विशेष महत्व, उपयोगिता और सकारात्मक प्रभाव की गहन समीक्षा।"},
                {"title": f"{s_title} - भावी दृष्टिकोण", "desc": f"आगामी समय में सतत विकास और व्यावहारिक सफलता हेतु निर्धारित रणनीतिक दिशा-निर्देश।"}
            ]

        num_cards = min(len(cards), 4)
        total_w = Inches(11.733)
        gap = Inches(0.25)
        card_w = (total_w - (gap * (num_cards - 1))) / num_cards
        card_h = Inches(4.9)
        top = Inches(1.75)

        for idx, card_data in enumerate(cards[:num_cards]):
            if isinstance(card_data, str):
                card_data = {"title": card_data, "desc": ""}
            elif not isinstance(card_data, dict):
                card_data = {"title": str(card_data), "desc": ""}
            left = Inches(0.8) + (idx * (card_w + gap))
            card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, card_w, card_h)
            card.fill.solid()
            card.fill.fore_color.rgb = theme["card_bg"]
            card.line.color.rgb = theme.get("card_border", theme["accent"])
            card.line.width = Pt(1.2)

            t_box = slide.shapes.add_textbox(left + Inches(0.25), top + Inches(0.25), card_w - Inches(0.5), card_h - Inches(0.5))
            tf = t_box.text_frame
            tf.word_wrap = True

            p_title = tf.paragraphs[0]
            p_title.text = (card_data.get("title") or f"Pillar {idx + 1}").upper()
            p_title.font.bold = True
            p_title.font.size = Pt(11)
            p_title.font.color.rgb = theme["accent"]
            p_title.space_after = Pt(12)

            stat = card_data.get("stat")
            if stat:
                p_stat = tf.add_paragraph()
                p_stat.text = str(stat)
                p_stat.font.bold = True
                p_stat.font.size = Pt(36)
                p_stat.font.color.rgb = theme["text_primary"]
                p_stat.space_after = Pt(14)

            p_desc = tf.add_paragraph()
            p_desc.text = card_data.get("desc", "")
            p_desc.font.size = Pt(12)
            p_desc.font.color.rgb = theme["text_secondary"]
            p_desc.line_spacing = 1.3

    # ==========================================
    # SLIDE BUILDER 4: NATIVE EXCEL CHART SLIDE
    # ==========================================
    def add_chart_slide(self, prs, theme: Dict[str, Any], data: Dict[str, Any], current_slide: int = 4, total_slides: int = 8, presenter: str = "Rhynia AI"):
        blank_slide_layout = prs.slide_layouts[6]
        slide = prs.slides.add_slide(blank_slide_layout)
        self._apply_slide_background(slide, theme)
        self._add_header(slide, theme, data.get("title", "Market Performance Analytics"), data.get("category", "DATA ANALYTICS"))
        self._add_footer(slide, theme, current_slide, total_slides, presenter)

        chart_data = CategoryChartData()
        chart_data.categories = data.get("categories", ["2023", "2024", "2025", "2026"])
        for s in data.get("series", [{"name": "Metric Index", "values": [35, 55, 80, 115]}]):
            chart_data.add_series(s.get("name", "Metric"), tuple(s.get("values", [10, 20, 30, 40])))

        ctype_str = (data.get("chart_type") or "column").lower()
        if ctype_str == "bar":
            c_type = XL_CHART_TYPE.BAR_CLUSTERED
        elif ctype_str in ["pie", "donut"]:
            c_type = XL_CHART_TYPE.PIE
        elif ctype_str == "line":
            c_type = XL_CHART_TYPE.LINE
        else:
            c_type = XL_CHART_TYPE.COLUMN_CLUSTERED

        # Native Chart Shape
        chart_shape = slide.shapes.add_chart(
            c_type, Inches(0.8), Inches(1.75), Inches(7.5), Inches(4.9), chart_data
        )
        chart = chart_shape.chart
        chart.has_legend = len(data.get("series", [])) > 1 or ctype_str in ["pie", "donut"]
        if chart.has_legend:
            chart.legend.position = XL_LEGEND_POSITION.TOP
            chart.legend.include_in_layout = False

        # Right Summary Card
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.6), Inches(1.75), Inches(3.933), Inches(4.9))
        card.fill.solid()
        card.fill.fore_color.rgb = theme["card_bg"]
        card.line.color.rgb = theme.get("card_border", theme["accent"])
        card.line.width = Pt(1)

        t_box = slide.shapes.add_textbox(Inches(8.85), Inches(1.95), Inches(3.433), Inches(4.5))
        tf = t_box.text_frame
        tf.word_wrap = True

        takeaways = data.get("takeaways", [])
        if not takeaways:
            s_title = data.get("title", "डेटा विश्लेषण")
            takeaways = [
                f"{s_title} के आंकड़ों में निरंतर और सकारात्मक प्रगति दर्ज की गई है।",
                f"वैज्ञानिक व सांख्यिकीय विश्लेषण से लक्ष्यों की प्राप्ति और उच्च दक्षता सिद्ध होती है।",
                f"आगामी समय में बेहतर परिणामों और रणनीतिक विकास हेतु निरंतर निगरानी आवश्यक है।"
            ]

        is_hi = any(ord(c) > 127 for c in str(takeaways))
        p_h = tf.paragraphs[0]
        p_h.text = "प्रमुख निष्कर्ष एवं अंतर्दृष्टि" if is_hi else "STRATEGIC TAKEAWAYS"
        p_h.font.bold = True
        p_h.font.size = Pt(12)
        p_h.font.color.rgb = theme["accent"]
        p_h.space_after = Pt(14)

        for pt in takeaways:
            if isinstance(pt, dict):
                pt_text = pt.get("desc") or pt.get("title") or " - ".join(str(v) for v in pt.values())
            else:
                pt_text = str(pt)
            p = tf.add_paragraph()
            p.text = f"✔  {pt_text}"
            p.font.size = Pt(11.5)
            p.font.color.rgb = theme["text_secondary"]
            p.space_after = Pt(10)
            p.line_spacing = 1.25

    # ==========================================
    # SLIDE BUILDER 5: COMPARATIVE DATA TABLE
    # ==========================================
    def add_table_slide(self, prs, theme: Dict[str, Any], data: Dict[str, Any], current_slide: int = 5, total_slides: int = 8, presenter: str = "Rhynia AI"):
        blank_slide_layout = prs.slide_layouts[6]
        slide = prs.slides.add_slide(blank_slide_layout)
        self._apply_slide_background(slide, theme)
        self._add_header(slide, theme, data.get("title", "Comparative Architecture Evaluation"), data.get("category", "COMPARATIVE MATRIX"))
        self._add_footer(slide, theme, current_slide, total_slides, presenter)

        s_title = data.get("title", "तुलनात्मक मूल्यांकन")
        headers = data.get("headers", [])
        if not headers:
            headers = ["मुख्य आयाम / घटक", "पारंपरिक / आधार स्थिति", "आधुनिक / अनुकूलित मॉडल", "प्रभाव व परिणाम"]

        rows = data.get("rows", [])
        if not rows:
            rows = [
                [f"{s_title} - आयाम 1", "सीमित संरचना", "सक्रिय वैज्ञानिक प्रबंधन", "उच्च दक्षता व सटीकता"],
                [f"{s_title} - आयाम 2", "पारंपरिक विधि", "उन्नत तकनीकी एकीकरण", "विश्वसनीय और तेज परिणाम"],
                [f"{s_title} - आयाम 3", "असंगठित प्रक्रिया", "स्वचालित अनुकूलन", "दीर्घकालिक स्थिरता व प्रभाव"]
            ]

        num_cols = len(headers)
        num_rows = len(rows) + 1
        t_left = Inches(0.8)
        t_top = Inches(1.8)
        t_width = Inches(11.733)
        t_height = Inches(4.8)

        table_shape = slide.shapes.add_table(num_rows, num_cols, t_left, t_top, t_width, t_height)
        table = table_shape.table

        # Header Row
        for c_idx, h_text in enumerate(headers):
            cell = table.cell(0, c_idx)
            cell.fill.solid()
            cell.fill.fore_color.rgb = theme["table_header"]
            cell.text_frame.margin_left = cell.text_frame.margin_right = Inches(0.15)
            p = cell.text_frame.paragraphs[0]
            p.text = str(h_text)
            p.font.bold = True
            p.font.size = Pt(12)
            p.font.color.rgb = RGBColor(255, 255, 255)
            p.alignment = PP_ALIGN.LEFT

        # Data Rows
        for r_idx, row in enumerate(rows):
            is_alt = (r_idx % 2 == 1)
            row_color = theme.get("table_row_alt", theme["card_bg"]) if is_alt else theme["card_bg"]

            if isinstance(row, dict):
                row_vals = list(row.values())
            elif isinstance(row, (list, tuple)):
                row_vals = list(row)
            else:
                row_vals = [str(row)]

            for c_idx in range(num_cols):
                raw_v = row_vals[c_idx] if c_idx < len(row_vals) else ""
                val_str = str(raw_v).strip() if raw_v is not None else ""
                if not val_str:
                    val_str = f"{s_title} घटक {r_idx+1}.{c_idx+1}"

                cell = table.cell(r_idx + 1, c_idx)
                cell.fill.solid()
                cell.fill.fore_color.rgb = row_color
                cell.text_frame.margin_left = cell.text_frame.margin_right = Inches(0.15)
                p = cell.text_frame.paragraphs[0]
                p.text = val_str
                p.font.size = Pt(11)
                p.font.color.rgb = theme["text_primary"]
                p.alignment = PP_ALIGN.LEFT

    # ==========================================
    # SLIDE BUILDER 6: PROCESS / WORKFLOW ROADMAP
    # ==========================================
    def add_process_diagram_slide(self, prs, theme: Dict[str, Any], data: Dict[str, Any], current_slide: int = 6, total_slides: int = 8, presenter: str = "Rhynia AI"):
        blank_slide_layout = prs.slide_layouts[6]
        slide = prs.slides.add_slide(blank_slide_layout)
        self._apply_slide_background(slide, theme)
        self._add_header(slide, theme, data.get("title", "Phased Execution Roadmap"), data.get("category", "WORKFLOW ROADMAP"))
        self._add_footer(slide, theme, current_slide, total_slides, presenter)

        steps = data.get("steps", [])
        if not steps:
            s_title = data.get("title", "कार्यप्रणाली")
            steps = [
                {"phase": "चरण 1", "title": "प्रारंभिक तैयारी एवं शोध", "desc": f"{s_title} की पृष्ठभूमि, प्राथमिक आवश्यकताओं और बुनियादी संरचना की विस्तृत तैयारी।"},
                {"phase": "चरण 2", "title": "कार्यान्वयन एवं परीक्षण", "desc": "योजना के अनुसार मुख्य प्रक्रियाओं का संचालन और वास्तविक परिस्थितियों में गहन परीक्षण।"},
                {"phase": "चरण 3", "title": "समीक्षा एवं अनुकूलन", "desc": "प्राप्त परिणामों का वैज्ञानिक विश्लेषण कर आवश्यक सुधार और दक्षता में वृद्धि करना।"},
                {"phase": "चरण 4", "title": "अंतिम लक्ष्य व सफलता", "desc": "दीर्घकालिक उद्देश्यों की पूर्ण प्राप्ति और सतत प्रगति हेतु मानक स्थापित करना।"}
            ]

        num_steps = min(len(steps), 5)
        total_w = Inches(11.733)
        arrow_w = Inches(0.4)
        box_w = (total_w - (arrow_w * (num_steps - 1))) / num_steps
        box_h = Inches(4.8)
        top = Inches(1.8)

        for idx, st in enumerate(steps[:num_steps]):
            if isinstance(st, str):
                st = {"phase": f"चरण {idx + 1}", "title": st, "desc": ""}
            elif not isinstance(st, dict):
                st = {"phase": f"चरण {idx + 1}", "title": str(st), "desc": ""}
            left = Inches(0.8) + (idx * (box_w + arrow_w))
            card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, box_w, box_h)
            card.fill.solid()
            card.fill.fore_color.rgb = theme["card_bg"]
            card.line.color.rgb = theme.get("card_border", theme["accent"])
            card.line.width = Pt(1.2)

            t_box = slide.shapes.add_textbox(left + Inches(0.18), top + Inches(0.2), box_w - Inches(0.36), box_h - Inches(0.4))
            tf = t_box.text_frame
            tf.word_wrap = True

            p_ph = tf.paragraphs[0]
            p_ph.text = st.get("phase", f"STEP {idx + 1}").upper()
            p_ph.font.bold = True
            p_ph.font.size = Pt(11)
            p_ph.font.color.rgb = theme["accent"]
            p_ph.space_after = Pt(12)

            p_t = tf.add_paragraph()
            p_t.text = st.get("title", f"Step {idx + 1}")
            p_t.font.bold = True
            p_t.font.size = Pt(15)
            p_t.font.color.rgb = theme["text_primary"]
            p_t.space_after = Pt(12)

            p_d = tf.add_paragraph()
            p_d.text = st.get("desc", "")
            p_d.font.size = Pt(11)
            p_d.font.color.rgb = theme["text_secondary"]
            p_d.line_spacing = 1.3

            if idx < num_steps - 1:
                arr_left = left + box_w
                arrow = slide.shapes.add_shape(
                    MSO_SHAPE.CHEVRON, arr_left + Inches(0.06), top + Inches(2.1), Inches(0.28), Inches(0.45)
                )
                arrow.fill.solid()
                arrow.fill.fore_color.rgb = theme["accent"]
                arrow.line.fill.background()

    # ==========================================
    # SLIDE BUILDER 7: SPLIT TWO-COLUMN COMPARISON
    # ==========================================
    def add_split_comparison_slide(self, prs, theme: Dict[str, Any], data: Dict[str, Any], current_slide: int = 7, total_slides: int = 8, presenter: str = "Rhynia AI"):
        blank_slide_layout = prs.slide_layouts[6]
        slide = prs.slides.add_slide(blank_slide_layout)
        self._apply_slide_background(slide, theme)
        self._add_header(slide, theme, data.get("title", "Comparative Analysis"), data.get("category", "COMPARISON"))
        self._add_footer(slide, theme, current_slide, total_slides, presenter)

        s_title = data.get("title", "तुलनात्मक विश्लेषण")
        left_col = data.get("left_column")
        if not left_col or not left_col.get("points"):
            left_col = {
                "title": "पारंपरिक दृष्टिकोण / पूर्व स्थिति",
                "points": [
                    f"{s_title} से जुड़ी पारंपरिक पद्धतियों में प्रक्रियात्मक विलंब और सीमित संसाधन दक्षता की समस्या।",
                    "डेटा और सूचनाओं के बिखराव के कारण त्वरित निर्णय लेने और विश्लेषण में कठिनाई।",
                    "सीमित तकनीकी एकीकरण और मैन्युअल निर्भरता के कारण उच्च त्रुटि दर और जोखिम।"
                ]
            }
        right_col = data.get("right_column")
        if not right_col or not right_col.get("points"):
            right_col = {
                "title": "आधुनिक वैज्ञानिक मॉडल / उन्नत स्थिति",
                "points": [
                    f"उन्नत तकनीकों और व्यवस्थित योजना द्वारा {s_title} की कार्यक्षमता में बहुगुणा सुधार।",
                    "वास्तविक समय में सटीक निगरानी और वैज्ञानिक विश्लेषण द्वारा त्रुटिहीन एवं ठोस परिणाम।",
                    "दीर्घकालिक स्थिरता, व्यापक मापनीयता और सतत विकास हेतु अनुकूलित संरचना।"
                ]
            }

        col_w = Inches(5.65)
        col_h = Inches(4.8)
        top = Inches(1.8)

        # Left Card
        left_card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), top, col_w, col_h)
        left_card.fill.solid()
        left_card.fill.fore_color.rgb = theme["card_bg"]
        left_card.line.color.rgb = theme.get("card_border", theme["accent"])
        left_card.line.width = Pt(1)

        t_left = slide.shapes.add_textbox(Inches(1.05), top + Inches(0.25), col_w - Inches(0.5), col_h - Inches(0.5))
        tf_l = t_left.text_frame
        tf_l.word_wrap = True
        p_lt = tf_l.paragraphs[0]
        p_lt.text = left_col.get("title", "Left Pillar").upper()
        p_lt.font.bold = True
        p_lt.font.size = Pt(13)
        p_lt.font.color.rgb = theme.get("accent3", theme["accent"])
        p_lt.space_after = Pt(14)

        for pt in left_col.get("points", []):
            pt_txt = f"{pt.get('title', '')}: {pt.get('desc', '')}".strip(": ") if isinstance(pt, dict) else str(pt)
            p = tf_l.add_paragraph()
            p.text = f"•  {pt_txt}"
            p.font.size = Pt(12)
            p.font.color.rgb = theme["text_secondary"]
            p.space_after = Pt(10)
            p.line_spacing = 1.3

        # Right Card
        right_card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.883), top, col_w, col_h)
        right_card.fill.solid()
        right_card.fill.fore_color.rgb = theme["card_bg"]
        right_card.line.color.rgb = theme["accent"]
        right_card.line.width = Pt(1.5)

        t_right = slide.shapes.add_textbox(Inches(7.133), top + Inches(0.25), col_w - Inches(0.5), col_h - Inches(0.5))
        tf_r = t_right.text_frame
        tf_r.word_wrap = True
        p_rt = tf_r.paragraphs[0]
        p_rt.text = right_col.get("title", "Right Pillar").upper()
        p_rt.font.bold = True
        p_rt.font.size = Pt(13)
        p_rt.font.color.rgb = theme["accent"]
        p_rt.space_after = Pt(14)

        for pt in right_col.get("points", []):
            pt_txt = f"{pt.get('title', '')}: {pt.get('desc', '')}".strip(": ") if isinstance(pt, dict) else str(pt)
            p = tf_r.add_paragraph()
            p.text = f"✔  {pt_txt}"
            p.font.size = Pt(12)
            p.font.color.rgb = theme["text_primary"]
            p.space_after = Pt(10)
            p.line_spacing = 1.3

    # ==========================================
    # SLIDE BUILDER 8: 4-BOX QUADRANT / MATRIX
    # ==========================================
    def add_quadrant_matrix_slide(self, prs, theme: Dict[str, Any], data: Dict[str, Any], current_slide: int = 8, total_slides: int = 8, presenter: str = "Rhynia AI"):
        blank_slide_layout = prs.slide_layouts[6]
        slide = prs.slides.add_slide(blank_slide_layout)
        self._apply_slide_background(slide, theme)
        self._add_header(slide, theme, data.get("title", "Core Strategic Quadrant"), data.get("category", "MATRIX FRAMEWORK"))
        self._add_footer(slide, theme, current_slide, total_slides, presenter)

        quadrants = data.get("quadrants", [])
        if not quadrants:
            s_title = data.get("title", "आयाम")
            quadrants = [
                {"title": "प्रमुख शक्तियां एवं क्षमताएं", "desc": f"{s_title} के अंतर्गत अंतर्निहित सामर्थ्य, सुदृढ़ आधार और प्राथमिक सफलता के निर्णायक कारक।"},
                {"title": "विकास एवं नए अवसर", "desc": f"आगामी समय में विस्तार की संभावनाएं, तकनीकी नवाचार और नए व्यावहारिक क्षेत्रों में व्यापक अनुप्रयोग।"},
                {"title": "चुनौतियां एवं जोखिम समाधान", "desc": f"संभावित बाधाओं का समयबद्ध वैज्ञानिक विश्लेषण, सुरक्षा मानक और प्रभावी समाधान रणनीतियां।"},
                {"title": "अंतिम लक्ष्य व उपलब्धियां", "desc": f"निर्धारित समयसीमा में प्रमुख मील के पत्थर, गुणवत्ता मानक और सतत दीर्घकालिक प्रभाव।"}
            ]

        col_w = Inches(5.65)
        row_h = Inches(2.25)
        top_start = Inches(1.8)

        coords = [
            (Inches(0.8), top_start),
            (Inches(6.883), top_start),
            (Inches(0.8), top_start + Inches(2.45)),
            (Inches(6.883), top_start + Inches(2.45))
        ]

        for idx, (left, top) in enumerate(coords):
            q_data = quadrants[idx] if idx < len(quadrants) else {"title": f"Quadrant {idx+1}", "desc": ""}
            if isinstance(q_data, str):
                q_data = {"title": q_data, "desc": ""}
            elif not isinstance(q_data, dict):
                q_data = {"title": str(q_data), "desc": ""}
            card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, col_w, row_h)
            card.fill.solid()
            card.fill.fore_color.rgb = theme["card_bg"]
            card.line.color.rgb = theme.get("card_border", theme["accent"])
            card.line.width = Pt(1)

            t_box = slide.shapes.add_textbox(left + Inches(0.25), top + Inches(0.2), col_w - Inches(0.5), row_h - Inches(0.4))
            tf = t_box.text_frame
            tf.word_wrap = True

            p_t = tf.paragraphs[0]
            p_t.text = f"0{idx+1}  {q_data.get('title', '').upper()}"
            p_t.font.bold = True
            p_t.font.size = Pt(13)
            p_t.font.color.rgb = theme["accent"]
            p_t.space_after = Pt(8)

            p_d = tf.add_paragraph()
            p_d.text = q_data.get("desc", "")
            p_d.font.size = Pt(11)
            p_d.font.color.rgb = theme["text_secondary"]
            p_d.line_spacing = 1.25

    # ==========================================
    # SLIDE BUILDER 9: IMAGE & CONTENT SHOWCASE
    # ==========================================
    def add_image_content_slide(self, prs, theme: Dict[str, Any], data: Dict[str, Any], current_slide: int = 7, total_slides: int = 8, presenter: str = "Rhynia AI"):
        blank_slide_layout = prs.slide_layouts[6]
        slide = prs.slides.add_slide(blank_slide_layout)
        self._apply_slide_background(slide, theme)
        self._add_header(slide, theme, data.get("title", "Visual Overview & Architecture"), data.get("category", "VISUAL GROUNDING"))
        self._add_footer(slide, theme, current_slide, total_slides, presenter)

        image_url = data.get("image_url")
        image_path = self._download_and_cache_image(image_url) if image_url else None

        # Left Points Card
        t_card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(6.3), Inches(4.8))
        t_card.fill.solid()
        t_card.fill.fore_color.rgb = theme["card_bg"]
        t_card.line.color.rgb = theme.get("card_border", theme["accent"])
        t_card.line.width = Pt(1)

        t_box = slide.shapes.add_textbox(Inches(1.05), Inches(2.0), Inches(5.8), Inches(4.4))
        tf = t_box.text_frame
        tf.word_wrap = True

        points = data.get("points", [])
        if not points:
            s_title = data.get("title", "दृश्य अवलोकन")
            points = [
                f"{s_title} से संबंधित प्रमुख संरचनात्मक एवं व्यावहारिक घटकों का स्पष्ट दृश्य निरूपण।",
                "प्रत्यक्ष प्रमाणों और वैज्ञानिक आंकड़ों के माध्यम से संपूर्ण प्रक्रिया की प्रामाणिक पुष्टि।",
                "सैद्धांतिक अवधारणाओं को व्यावहारिक संदर्भ में समझने हेतु विस्तृत और प्रभावी विश्लेषण।"
            ]

        is_hi = any(ord(c) > 127 for c in str(points))
        p_h = tf.paragraphs[0]
        p_h.text = (data.get("subtitle") or ("प्रमुख अवलोकन एवं बिंदु" if is_hi else "KEY OBSERVATIONS & HIGHLIGHTS")).upper()
        p_h.font.bold = True
        p_h.font.size = Pt(12)
        p_h.font.color.rgb = theme["accent"]
        p_h.space_after = Pt(14)

        for pt in points:
            pt_txt = f"{pt.get('title', '')}: {pt.get('desc', '')}".strip(": ") if isinstance(pt, dict) else str(pt)
            p = tf.add_paragraph()
            p.text = f"✔  {pt_txt}"
            p.font.size = Pt(12)
            p.font.color.rgb = theme["text_secondary"]
            p.space_after = Pt(10)
            p.line_spacing = 1.3

        # Right Image Frame
        img_frame = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.45), Inches(1.8), Inches(5.083), Inches(4.8))
        img_frame.fill.solid()
        img_frame.fill.fore_color.rgb = theme["card_bg"]
        img_frame.line.color.rgb = theme["accent"]
        img_frame.line.width = Pt(1.5)

        if image_path:
            try:
                slide.shapes.add_picture(str(image_path), Inches(7.55), Inches(1.9), Inches(4.883), Inches(4.6))
            except Exception as e:
                logger.warning(f"Error adding picture: {e}")
        else:
            tf_img = img_frame.text_frame
            p_ph = tf_img.paragraphs[0]
            p_ph.text = data.get("image_caption", "Topic Visualization")
            p_ph.font.size = Pt(13)
            p_ph.font.color.rgb = theme["text_secondary"]
            p_ph.alignment = PP_ALIGN.CENTER

    # ==========================================
    # SLIDE BUILDER 10: CONCLUSION & NEXT STEPS
    # ==========================================
    def add_conclusion_slide(self, prs, theme: Dict[str, Any], data: Dict[str, Any], current_slide: int = 8, total_slides: int = 8, presenter: str = "Rhynia AI"):
        blank_slide_layout = prs.slide_layouts[6]
        slide = prs.slides.add_slide(blank_slide_layout)
        self._apply_slide_background(slide, theme)
        self._add_header(slide, theme, data.get("title", "Strategic Summary & Next Steps"), data.get("category", "EXECUTIVE SUMMARY"))
        self._add_footer(slide, theme, current_slide, total_slides, presenter)

        # Left Checklist Card
        takeaways = data.get("takeaways", [])
        if not takeaways:
            s_title = data.get("title", "निष्कर्ष")
            takeaways = [
                f"{s_title} के संपूर्ण विश्लेषण से स्पष्ट होता है कि ठोस योजना और वैज्ञानिक दृष्टिकोण से सर्वोच्च परिणाम प्राप्त होते हैं।",
                "आधुनिक तकनीकों और प्रामाणिक विधियों के समन्वय से सभी प्रमुख लक्ष्यों की समयबद्ध और सुरक्षित प्राप्ति संभव है।",
                "सतत नवाचार, नियमित मूल्यांकन और सहयोगात्मक प्रयासों द्वारा भविष्य में दीर्घकालिक सफलता सुनिश्चित होती है।"
            ]

        is_hi = any(ord(c) > 127 for c in str(takeaways))

        card_l = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(7.5), Inches(4.8))
        card_l.fill.solid()
        card_l.fill.fore_color.rgb = theme["card_bg"]
        card_l.line.color.rgb = theme.get("card_border", theme["accent"])
        card_l.line.width = Pt(1)

        t_box_l = slide.shapes.add_textbox(Inches(1.1), Inches(2.05), Inches(6.9), Inches(4.3))
        tf_l = t_box_l.text_frame
        tf_l.word_wrap = True

        p_head = tf_l.paragraphs[0]
        p_head.text = "प्रमुख निष्कर्ष एवं परिणाम" if is_hi else "CORE STRATEGIC DELIVERABLES"
        p_head.font.bold = True
        p_head.font.size = Pt(13)
        p_head.font.color.rgb = theme["accent"]
        p_head.space_after = Pt(14)

        for t in takeaways:
            t_txt = f"{t.get('title', '')}: {t.get('desc', '')}".strip(": ") if isinstance(t, dict) else str(t)
            p_bullet = tf_l.add_paragraph()
            p_bullet.text = f"✔   {t_txt}"
            p_bullet.font.size = Pt(12.5)
            p_bullet.font.color.rgb = theme["text_primary"]
            p_bullet.space_after = Pt(14)
            p_bullet.line_spacing = 1.3

        # Right Contact / Action Box
        card_r = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.65), Inches(1.8), Inches(3.883), Inches(4.8))
        card_r.fill.solid()
        card_r.fill.fore_color.rgb = theme["card_bg"]
        card_r.line.color.rgb = theme["accent"]
        card_r.line.width = Pt(1.5)

        t_box_r = slide.shapes.add_textbox(Inches(8.9), Inches(2.05), Inches(3.383), Inches(4.3))
        tf_r = t_box_r.text_frame
        tf_r.word_wrap = True

        p_qa = tf_r.paragraphs[0]
        p_qa.text = "धन्यवाद" if is_hi else "THANK YOU"
        p_qa.font.bold = True
        p_qa.font.size = Pt(24)
        p_qa.font.color.rgb = theme["text_primary"]
        p_qa.space_after = Pt(6)

        p_qa_sub = tf_r.add_paragraph()
        p_qa_sub.text = "प्रश्नोत्तरी एवं परिचर्चा" if is_hi else "Questions & Collaborative Discussion"
        p_qa_sub.font.size = Pt(13)
        p_qa_sub.font.color.rgb = theme["accent"]
        p_qa_sub.space_after = Pt(16)

        p_contact = tf_r.add_paragraph()
        p_contact.text = data.get("contact_info", f"Presented by: {presenter}\nRhynia AI Platform\nhttps://rhynia.ai")
        p_contact.font.size = Pt(11)
        p_contact.font.color.rgb = theme["text_secondary"]
        p_contact.line_spacing = 1.3

    # ==========================================
    # PRESENTATION COMPILER
    # ==========================================
    def build_presentation(self, presentation_data: Dict[str, Any], user_id: str = "guest") -> Dict[str, Any]:
        """
        Compiles presentation JSON into a native .pptx file and saves to disk.
        """
        prs = Presentation()
        prs.slide_width = Inches(SLIDE_WIDTH_INCHES)
        prs.slide_height = Inches(SLIDE_HEIGHT_INCHES)

        theme_name = presentation_data.get("theme", "executive_dark")
        theme = THEMES.get(theme_name, THEMES["executive_dark"])
        presenter = presentation_data.get("presenter", "Rhynia AI")

        slides_spec = presentation_data.get("slides", [])
        if not slides_spec:
            slides_spec = [
                {"layout": "title", "title": presentation_data.get("title", "Presentation"), "subtitle": "Executive Overview"},
                {"layout": "cards", "title": "Strategic Focus", "cards": []},
                {"layout": "conclusion", "title": "Conclusion"}
            ]

        total_slides = len(slides_spec)

        for idx, s in enumerate(slides_spec):
            curr_slide = idx + 1
            layout_type = (s.get("layout") or "cards").lower()
            try:
                if layout_type == "title":
                    self.add_title_slide(prs, theme, s)
                elif layout_type == "agenda":
                    self.add_agenda_slide(prs, theme, s, curr_slide, total_slides, presenter)
                elif layout_type in ["cards", "content", "bullets", "kpi"]:
                    self.add_card_content_slide(prs, theme, s, curr_slide, total_slides, presenter)
                elif layout_type == "table":
                    self.add_table_slide(prs, theme, s, curr_slide, total_slides, presenter)
                elif layout_type in ["chart", "graph"]:
                    self.add_chart_slide(prs, theme, s, curr_slide, total_slides, presenter)
                elif layout_type in ["process", "flowchart", "timeline", "roadmap"]:
                    self.add_process_diagram_slide(prs, theme, s, curr_slide, total_slides, presenter)
                elif layout_type in ["split", "comparison", "two_column"]:
                    self.add_split_comparison_slide(prs, theme, s, curr_slide, total_slides, presenter)
                elif layout_type in ["matrix", "quadrant", "grid_4"]:
                    self.add_quadrant_matrix_slide(prs, theme, s, curr_slide, total_slides, presenter)
                elif layout_type in ["image_content", "showcase", "image"]:
                    self.add_image_content_slide(prs, theme, s, curr_slide, total_slides, presenter)
                elif layout_type == "conclusion":
                    self.add_conclusion_slide(prs, theme, s, curr_slide, total_slides, presenter)
                else:
                    self.add_card_content_slide(prs, theme, s, curr_slide, total_slides, presenter)
            except Exception as e:
                logger.error(f"Error building slide layout {layout_type}: {e}")
                self.add_card_content_slide(prs, theme, s, curr_slide, total_slides, presenter)

        # Output target directory
        user_dir = self.output_dir / str(user_id)
        user_dir.mkdir(parents=True, exist_ok=True)

        file_id = uuid.uuid4().hex[:12]
        clean_title = "".join(c for c in presentation_data.get("title", "Rhynia_Presentation") if c.isalnum() or c in (" ", "_", "-")).strip()
        clean_title = clean_title.replace(" ", "_")[:40] or "Presentation"
        filename = f"{clean_title}_{file_id}.pptx"
        output_path = user_dir / filename

        prs.save(str(output_path))
        file_size = output_path.stat().st_size

        return {
            "file_id": file_id,
            "filename": filename,
            "file_path": str(output_path),
            "file_size": file_size,
            "slide_count": len(prs.slides),
            "theme": theme_name,
            "title": presentation_data.get("title", "Presentation")
        }

    def build_word_document(self, presentation_data: Dict[str, Any], user_id: Optional[str] = None) -> Dict[str, Any]:
        """
        Builds a comprehensive Microsoft Word document (.docx) from presentation specification.
        Includes title, table of contents, executive summary, tables, and bullet points.
        """
        import docx
        from docx.shared import Pt as DocxPt, RGBColor as DocxRGBColor
        from docx.enum.text import WD_ALIGN_PARAGRAPH

        doc = docx.Document()
        title = presentation_data.get("title", "Executive Presentation Report")
        presenter = presentation_data.get("presenter", "Rhynia AI")
        slides = presentation_data.get("slides", [])

        # Document Title
        p_title = doc.add_paragraph()
        p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run_title = p_title.add_run(title)
        run_title.bold = True
        run_title.font.size = DocxPt(24)
        run_title.font.color.rgb = DocxRGBColor(0, 120, 212)

        # Subtitle
        p_sub = doc.add_paragraph()
        p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run_sub = p_sub.add_run(f"Executive Report & Briefing Document\nPresented by: {presenter} | Rhynia Intelligence Platform")
        run_sub.font.size = DocxPt(11)
        run_sub.font.color.rgb = DocxRGBColor(100, 116, 139)

        # Slide 1 Executive Metadata Box
        if slides and (slides[0].get("layout") == "title" or True):
            s1 = slides[0]
            p_box = doc.add_paragraph()
            p_box.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r_box = p_box.add_run("─── EXECUTIVE PRESENTATION METADATA ───")
            r_box.bold = True
            r_box.font.size = DocxPt(10)
            r_box.font.color.rgb = DocxRGBColor(0, 120, 212)

            t_meta = doc.add_table(rows=2, cols=2)
            t_meta.style = 'Table Grid'
            t_meta.cell(0, 0).paragraphs[0].text = f"Presented By: {s1.get('presented_by') or presenter}"
            t_meta.cell(0, 1).paragraphs[0].text = f"Date: {s1.get('date', '2026 Edition')}"
            t_meta.cell(1, 0).paragraphs[0].text = f"Location: {s1.get('location', 'Corporate HQ / New Delhi')}"
            t_meta.cell(1, 1).paragraphs[0].text = f"Goal: {s1.get('goal', 'Strategic Implementation & Action')}"

        doc.add_paragraph()

        for idx, slide in enumerate(slides, 1):
            s_title = slide.get("title", f"Slide {idx}")
            s_cat = slide.get("category", "EXECUTIVE BRIEFING")

            h = doc.add_heading(level=1)
            r_cat = h.add_run(f"[{s_cat.upper()}] ")
            r_cat.font.size = DocxPt(11)
            r_cat.font.color.rgb = DocxRGBColor(0, 120, 212)
            r_title = h.add_run(s_title)
            r_title.bold = True
            r_title.font.size = DocxPt(16)

            if slide.get("subtitle"):
                p = doc.add_paragraph()
                r = p.add_run(slide["subtitle"])
                r.italic = True
                r.font.size = DocxPt(10.5)

            if slide.get("items"):
                for itm in slide["items"]:
                    p = doc.add_paragraph(style='List Bullet')
                    r_b = p.add_run(itm.get("title", "") + ": ")
                    r_b.bold = True
                    p.add_run(itm.get("desc", ""))

            if slide.get("cards"):
                for c in slide["cards"]:
                    p = doc.add_paragraph(style='List Bullet')
                    if c.get("stat"):
                        r_stat = p.add_run(f"[{c['stat']}] ")
                        r_stat.bold = True
                    r_c = p.add_run(c.get("title", "") + " — ")
                    r_c.bold = True
                    p.add_run(c.get("desc", ""))

            if slide.get("takeaways"):
                for t in slide["takeaways"]:
                    p = doc.add_paragraph(style='List Bullet')
                    p.add_run(t)

            if slide.get("headers") and slide.get("rows"):
                headers = slide["headers"]
                rows = slide["rows"]
                table = doc.add_table(rows=len(rows) + 1, cols=len(headers))
                table.style = 'Table Grid'
                for c_idx, h_text in enumerate(headers):
                    table.cell(0, c_idx).text = str(h_text)
                for r_idx, row_vals in enumerate(rows):
                    for c_idx, cell_val in enumerate(row_vals):
                        table.cell(r_idx + 1, c_idx).text = str(cell_val)
                doc.add_paragraph()

            if slide.get("left_column") and slide.get("right_column"):
                l_col = slide["left_column"]
                r_col = slide["right_column"]
                p_l = doc.add_paragraph()
                r_l = p_l.add_run(f"• {l_col.get('title', 'Aspect 1')}:\n")
                r_l.bold = True
                for pt in l_col.get("points", []):
                    doc.add_paragraph(pt, style='List Bullet')

                p_r = doc.add_paragraph()
                r_r = p_r.add_run(f"• {r_col.get('title', 'Aspect 2')}:\n")
                r_r.bold = True
                for pt in r_col.get("points", []):
                    doc.add_paragraph(pt, style='List Bullet')

            if slide.get("steps"):
                for st in slide["steps"]:
                    p = doc.add_paragraph(style='List Number')
                    r_st = p.add_run(f"{st.get('phase', 'Step')}: {st.get('title', '')} — ")
                    r_st.bold = True
                    p.add_run(st.get("desc", ""))

        user_dir = self.output_dir / str(user_id)
        user_dir.mkdir(parents=True, exist_ok=True)
        file_id = uuid.uuid4().hex[:12]
        clean_title = "".join(c for c in title if c.isalnum() or c in (" ", "_", "-")).strip()
        clean_title = clean_title.replace(" ", "_")[:40] or "Report"
        doc_filename = f"{clean_title}_{file_id}.docx"
        doc_path = user_dir / doc_filename
        doc.save(str(doc_path))

        return {
            "file_id": file_id,
            "filename": doc_filename,
            "file_path": str(doc_path.resolve()),
            "file_size": doc_path.stat().st_size,
            "title": title
        }

    def build_pdf_document(self, presentation_data: Dict[str, Any], user_id: Optional[str] = None) -> Dict[str, Any]:
        """
        Builds a clean landscape presentation PDF document from presentation specification.
        """
        from reportlab.lib.pagesizes import letter, landscape
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.lib import colors

        title = presentation_data.get("title", "Executive Presentation")
        presenter = presentation_data.get("presenter", "Rhynia AI")
        slides = presentation_data.get("slides", [])

        user_dir = self.output_dir / str(user_id)
        user_dir.mkdir(parents=True, exist_ok=True)
        file_id = uuid.uuid4().hex[:12]
        clean_title = "".join(c for c in title if c.isalnum() or c in (" ", "_", "-")).strip()
        clean_title = clean_title.replace(" ", "_")[:40] or "Presentation"
        pdf_filename = f"{clean_title}_{file_id}.pdf"
        pdf_path = user_dir / pdf_filename

        doc = SimpleDocTemplate(
            str(pdf_path),
            pagesize=landscape(letter),
            rightMargin=36, leftMargin=36, topMargin=36, bottomMargin=36
        )

        styles = getSampleStyleSheet()
        title_style = ParagraphStyle(
            'DocTitle',
            parent=styles['Heading1'],
            fontSize=22,
            leading=26,
            textColor=colors.HexColor("#0078D4"),
            spaceAfter=10
        )
        cat_style = ParagraphStyle(
            'CatStyle',
            fontSize=10,
            leading=12,
            textColor=colors.HexColor("#0284C7"),
            spaceAfter=4
        )
        body_style = ParagraphStyle(
            'Body',
            parent=styles['Normal'],
            fontSize=11,
            leading=15,
            textColor=colors.HexColor("#1E293B"),
            spaceAfter=6
        )
        bullet_style = ParagraphStyle(
            'Bullet',
            parent=styles['Normal'],
            fontSize=10.5,
            leading=14,
            leftIndent=15,
            textColor=colors.HexColor("#334155"),
            spaceAfter=5
        )

        story = []

        for idx, slide in enumerate(slides):
            if idx > 0:
                story.append(PageBreak())

            cat = slide.get("category", "EXECUTIVE STRATEGY").upper()
            s_title = slide.get("title", "Overview")

            story.append(Paragraph(f"<b>❖  {cat}</b>", cat_style))
            story.append(Paragraph(s_title, title_style))
            if slide.get("subtitle"):
                story.append(Paragraph(f"<i>{slide['subtitle']}</i>", body_style))
            story.append(Spacer(1, 8))

            if idx == 0 or slide.get("layout") == "title":
                p_by = slide.get("presented_by") or presenter
                p_dt = slide.get("date", "2026 Edition")
                p_loc = slide.get("location", "Corporate HQ / New Delhi")
                p_gl = slide.get("goal", "Strategic Implementation & Action")
                meta_table_data = [
                    [Paragraph(f"<b>Presented By:</b> {p_by}", body_style), Paragraph(f"<b>Date:</b> {p_dt}", body_style)],
                    [Paragraph(f"<b>Location:</b> {p_loc}", body_style), Paragraph(f"<b>Goal:</b> {p_gl}", body_style)]
                ]
                t_m = Table(meta_table_data, colWidths=[260, 260])
                t_m.setStyle(TableStyle([
                    ('BACKGROUND', (0,0), (-1,-1), colors.HexColor('#F8FAFC')),
                    ('BOX', (0,0), (-1,-1), 1, colors.HexColor('#CBD5E1')),
                    ('INNERGRID', (0,0), (-1,-1), 0.5, colors.HexColor('#E2E8F0')),
                    ('TOPPADDING', (0,0), (-1,-1), 5),
                    ('BOTTOMPADDING', (0,0), (-1,-1), 5),
                ]))
                story.append(t_m)
                story.append(Spacer(1, 10))

            if slide.get("items"):
                for itm in slide["items"]:
                    story.append(Paragraph(f"• <b>{itm.get('title', '')}:</b> {itm.get('desc', '')}", bullet_style))

            if slide.get("cards"):
                for c in slide["cards"]:
                    stat_str = f"<b>[{c.get('stat')}]</b> " if c.get('stat') else ""
                    story.append(Paragraph(f"• {stat_str}<b>{c.get('title', '')}</b>: {c.get('desc', '')}", bullet_style))

            if slide.get("takeaways"):
                for t in slide["takeaways"]:
                    story.append(Paragraph(f"✔  {t}", bullet_style))

            if slide.get("steps"):
                for st in slide["steps"]:
                    story.append(Paragraph(f"<b>{st.get('phase', '')} - {st.get('title', '')}:</b> {st.get('desc', '')}", bullet_style))

            if slide.get("headers") and slide.get("rows"):
                table_data = [slide["headers"]] + slide["rows"]
                t = Table(table_data)
                t.setStyle(TableStyle([
                    ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#0078D4')),
                    ('TEXTCOLOR', (0,0), (-1,0), colors.whitesmoke),
                    ('FONTNAME', (0,0), (-1,0), 'Helvetica-Bold'),
                    ('FONTSIZE', (0,0), (-1,0), 9),
                    ('BOTTOMPADDING', (0,0), (-1,0), 5),
                    ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor('#CBD5E1')),
                ]))
                story.append(t)
                story.append(Spacer(1, 8))

            story.append(Spacer(1, 12))
            story.append(Paragraph(f"<font color='#64748B' size=8>Presented by: {presenter} | Rhynia AI • Slide {idx+1}/{len(slides)}</font>", body_style))

        try:
            doc.build(story)
        except Exception as pdf_err:
            logger.warning(f"Reportlab build warning: {pdf_err}")

        file_size = pdf_path.stat().st_size if pdf_path.exists() else 0
        return {
            "file_id": file_id,
            "filename": pdf_filename,
            "file_path": str(pdf_path.resolve()),
            "file_size": file_size,
            "title": title
        }


# Global Singleton
ppt_engine = PPTEngine()
